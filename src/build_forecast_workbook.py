"""Year-by-year outcomes of the owner-occupied-homes model, 2025-2035, every scenario:
outputs/Owner_Occupied_Forecast.xlsx. Values are the model's exact outputs (unrounded; cells are
formatted, not rounded). Run tenure_model.py first. Gridlines hidden, no freeze panes."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import tenure_model as tm

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "processed"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "outputs" / "Owner_Occupied_Forecast.xlsx"
FACE, INK2 = "Aptos", "52514E"
HDR = PatternFill("solid", fgColor="F0EFEC")
BASE = PatternFill("solid", fgColor="DCE8F7")
HAIR = Side(style="thin", color="E1E0D9")
YEARS = list(range(2025, 2036))
GROUPS = ["<25", "25-34", "35-44", "45-54", "55-64", "65-74", "75+"]

# model scenario name -> plain label, in display order
SCEN = {
    "Market path": "Mortgage from market rate path: BASE",
    "Fed dots path": "Mortgage from Fed dots path",
    "PMMS 6.0%": "Mortgage 6.0% flat",
    "PMMS 6.5%": "Mortgage 6.5% flat",
    "PMMS 7.0%": "Mortgage 7.0% flat (today's level)",
    "PMMS 7.5%": "Mortgage 7.5% flat",
    "PMMS 8.0%": "Mortgage 8.0% flat",
    "Market path, price/income flat at 2025 (no relief)": "Market path, price/income stays at 2025 (no relief)",
    "PMMS 8.0%, price/income flat at 2025 (no relief)": "8.0% flat, price/income stays at 2025 (no relief)",
    "Market path, delinquency back to 1991-2006 norm": "Market path, delinquency back to 1991-2006 norm",
    "Market path, model estimated on 2008-2025 only": "Market path, model estimated on 2008-2025 only",
    "Market path, 2024-25 young shortfall persists": "Market path, 2024-25 young shortfall persists",
    "Market path, main-immigration households": "Market path, JCHS main-immigration households",
    "Market path, zero-immigration households": "Market path, zero-immigration households",
    "Market path, households not calibrated to JCHS": "Market path, households not calibrated to JCHS",
    "rates by age frozen at 2025 (JCHS base-style)": "Ownership rates by age frozen at 2025 (no model)",
}
BASE_KEY = "Market path"


def f(bold=False, size=10, color="0B0B0B", italic=False):
    return Font(name=FACE, bold=bold, size=size, color=color, italic=italic)


def sheet(wb, name, title, note, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A1"].font = title, f(True, 14)
    ws["A2"], ws["A2"].font = note, f(size=9, color=INK2, italic=True)
    return ws


def table(ws, r, name, headers, rows, fmts, base_rows=()):
    ws.cell(r, 1, name).font = f(True, 11)
    for j, h in enumerate(headers, start=1):
        c = ws.cell(r + 1, j, h)
        c.font, c.fill, c.border = f(True, 9, INK2), HDR, Border(bottom=HAIR)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if j == 1 else "center")
    for i, row in enumerate(rows):
        for j, v in enumerate(row, start=1):
            c = ws.cell(r + 2 + i, j, None if (isinstance(v, float) and pd.isna(v)) else v)
            c.font = f(bold=i in base_rows)
            if i in base_rows:
                c.fill = BASE
            fm = fmts[j - 1] if isinstance(fmts, list) else fmts
            if fm and j > 1:
                c.number_format = fm
    return r + 2 + len(rows) + 1


def widths(ws, first, rest, n):
    ws.column_dimensions["A"].width = first
    for j in range(2, n + 2):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = rest


def note(ws, r, text):
    ws.cell(r, 1, text).font = f(size=9, color=INK2)
    return r + 1


# ------------------------------------------------------------------ model outputs
h, h_obs, hh, a, pop, w = tm.load()
p = tm.flows(h, w)
fit = tm.Fit(p, a)
fc, paths, assum = tm.forecast(a, h, hh, pop, w, fit, p)
fan = pd.read_csv(DATA / "owner_occupied_fan_base.csv", index_col=0)
fan.index = fan.index.astype(int)
keys = [k for k in SCEN if k in set(fc.scenario)]
base_i = keys.index(BASE_KEY)
piv = lambda col: fc.pivot_table(index="scenario", columns="year", values=col).loc[keys, YEARS]

wb = Workbook()

# 1. owner-occupied homes
ws = sheet(wb, "1 Owner-occupied homes", "US owner-occupied homes, model outcome by year (millions)",
           "Owners = households by age x ownership rate by age, summed. 2025 is actual (Census CPS/HVS Table 12a). "
           "BASE row shaded.", first=True)
own = piv("owners_k") / 1000
rows = [[SCEN[k]] + own.loc[k].tolist() + [(own.loc[k, 2031] - own.loc[k, 2025]) * 1000 / 6,
                                          (own.loc[k, 2035] - own.loc[k, 2025]) * 1000 / 10] for k in keys]
r = table(ws, 4, "A. Owner-occupied homes (m)", ["Scenario"] + YEARS + ["k a year 2025-31", "k a year 2025-35"], rows,
          ["0.000"] * 12 + ["#,##0", "#,##0"], base_rows=(base_i,))
chg = own.diff(axis=1) * 1000
rows = [[SCEN[k]] + chg.loc[k, YEARS[1:]].tolist() for k in keys]
r = table(ws, r + 1, "B. Change on the year (thousand homes)", ["Scenario"] + YEARS[1:], rows, "#,##0",
          base_rows=(base_i,))
lo = own.loc[BASE_KEY]
rows = [["10th percentile"] + [None] + fan.loc[YEARS[1:], "0.1"].tolist(),
        ["Median"] + [lo[2025]] + fan.loc[YEARS[1:], "0.5"].tolist(),
        ["90th percentile"] + [None] + fan.loc[YEARS[1:], "0.9"].tolist(),
        ["Deterministic base (row A)"] + lo.tolist()]
r = table(ws, r + 1, "C. Uncertainty around the BASE (m): 2,000 simulations", ["Band"] + YEARS, rows, "0.000")
r = note(ws, r, "Simulations draw the coefficients from their year-clustered covariance and add a whole historical "
                "year's residuals (all ages together) each year. Seed 7.")
widths(ws, 52, 10, 13)
ws.column_dimensions["M"].width = ws.column_dimensions["N"].width = 17

# 2. national rate and households
ws = sheet(wb, "2 Rate & households", "National ownership rate (%) and households (m), by year",
           "Ownership rate = owners / households. Households come from Census 2023 population projections by age, "
           "calibrated to JCHS 2025-35 totals.")
rt = piv("rate_pct")
r = table(ws, 4, "A. National homeownership rate (%)", ["Scenario"] + YEARS,
          [[SCEN[k]] + rt.loc[k].tolist() for k in keys], "0.00", base_rows=(base_i,))
hs = piv("households_k") / 1000
hs_rows, seen = [], set()
for k in keys:
    ser = fc[fc.scenario == k].households_series.iloc[0] + (" uncalibrated" if "not calibrated" in k else "")
    if ser in seen:
        continue
    seen.add(ser)
    label = {"low": "Low immigration, JCHS-calibrated (BASE)", "mid": "JCHS main immigration",
             "zero": "Zero immigration", "low uncalibrated": "Low immigration, not calibrated to JCHS"}[ser]
    hs_rows.append([label] + hs.loc[k].tolist())
r = table(ws, r + 1, "B. Households (m)", ["Household series"] + YEARS, hs_rows, "0.000", base_rows=(0,))
widths(ws, 52, 10, 12)

# 3. BASE by age group
ws = sheet(wb, "3 Base by age", "BASE scenario by age of householder, by year",
           "Mortgage rate from the market rate path, low-immigration households. Under-25 ownership held at its 2025 rate.")
path, H = paths[BASE_KEY]
grp = {g: [bd for bd in tm.BANDS if tm.GROUP[tm.SHORT[bd]] == g] for g in GROUPS}


def by_group(path, H):
    out = {}
    for yr in YEARS:
        hr = h.loc[2025] if yr == 2025 else path.loc[yr]
        wr = hh.loc[2025] if yr == 2025 else H.loc[yr]
        for g, bs in grp.items():
            own_k = float((wr[bs] * hr[bs] / 100).sum())
            out[(g, yr)] = (float(wr[bs].sum()), own_k, 100 * own_k / float(wr[bs].sum()))
    return out


bg = by_group(path, H)
r = table(ws, 4, "A. Ownership rate (%)", ["Age group"] + YEARS, [[g] + [bg[(g, y)][2] for y in YEARS] for g in GROUPS],
          "0.00")
r = table(ws, r + 1, "B. Owner-occupied homes (thousand)", ["Age group"] + YEARS,
          [[g] + [bg[(g, y)][1] for y in YEARS] for g in GROUPS] + [["Total"] + (own.loc[BASE_KEY] * 1000).tolist()],
          "#,##0", base_rows=(len(GROUPS),))
r = table(ws, r + 1, "C. Households (thousand)", ["Age group"] + YEARS,
          [[g] + [bg[(g, y)][0] for y in YEARS] for g in GROUPS] + [["Total"] + (hs.loc[BASE_KEY] * 1000).tolist()],
          "#,##0", base_rows=(len(GROUPS),))
widths(ws, 22, 10, 12)

# 4. every scenario by age group (long table)
ws = sheet(wb, "4 All scenarios by age", "Every scenario by age group and year (long table)",
           "One row per scenario, year and age group.")
rows = []
for k in keys:
    bgk = by_group(*paths[k])
    for yr in YEARS:
        for g in GROUPS:
            hh_k, own_k, rate = bgk[(g, yr)]
            rows.append([SCEN[k], yr, g, hh_k, own_k, rate])
r = table(ws, 4, "Scenario x year x age group", ["Scenario", "Year", "Age group", "Households (k)", "Owners (k)",
                                                 "Ownership rate (%)"], rows, [None, "0", None, "#,##0", "#,##0", "0.00"])
ws.auto_filter.ref = f"A5:F{4 + 1 + len(rows)}"
widths(ws, 52, 14, 6)

# 5. drivers
ws = sheet(wb, "5 Drivers", "Inputs by scenario and year",
           "The model reads LAST year's value of each driver, so a 2027 input moves the 2028 outcome.")
rows = []
for k in keys:
    pi = tm.price_income_path(a, YEARS[1:], "flat" if "flat" in k else "fannie")
    s = fc[fc.scenario == k].set_index("year")
    for yr in YEARS:
        rows.append([SCEN[k], yr, s.at[yr, "pmms"] if yr > 2025 else a.at[2025, "pmms_30y"],
                     a.at[2025, "price_to_income"] if yr == 2025 else pi[yr], s.at[yr, "pti_pct"], s.at[yr, "delinquency"]])
r = table(ws, 4, "Drivers (2025 = actual)", ["Scenario", "Year", "30y mortgage rate, PMMS (%)", "Price / income (NAR)",
                                             "Payment-to-income (%)", "Mortgage delinquency (%)"], rows,
          [None, "0", "0.00", "0.000", "0.00", "0.00"])
ws.auto_filter.ref = f"A5:F{4 + 1 + len(rows)}"
r = note(ws, r, "2026 mortgage rate = actual weekly average to 24-Sep-2026 (6.41%) blended with the scenario for the "
                "remaining weeks. The 'frozen' scenario uses no drivers.")
widths(ws, 52, 16, 6)

# 6. rates path
import mortgage_path as mp
mpa, _, mpi = mp.build("market")
mpd, _, mpdi = mp.build("dots")
kw = mpi["kim_wright"]
ws = sheet(wb, "6 Rates path", "BASE rates: 30y and 10y Treasury and the mortgage rate, from market expectations",
           f"As of {mpi['as_of']}: 30y Treasury {mpi['treasury_30y_now']:.2f}%, 10y Treasury "
           f"{mpi['treasury_10y_now']:.2f}%, effective fed funds {mpi['fed_funds_now']:.2f}%. Kim-Wright fit of "
           f"{mpi['kim_wright_date']}.")
rows = [[yr, r.expected_short_rate_avg,
         r.expected_30y_compounded, r.maturity_premium_30y, r.treasury_30y,
         r.expected_10y_compounded, r.maturity_premium_10y, r.treasury_10y, r.spread, r.mortgage_30y]
        for yr, r in mpa.iterrows()]
r = table(ws, 4, "BASE annual path: market expectations (year averages of the monthly path)",
          ["Year", "Expected short rate, year average %",
           "30y: compounded expected rate %", "30y: maturity premium (pp)", "30y Treasury %",
           "10y: compounded expected rate %", "10y: maturity premium (pp)", "10y Treasury %",
           "Mortgage - 30y spread (median, pp)", "30y mortgage (PMMS) %"], rows,
          [None, "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00"])
y30 = mpi["y1_to_y30_now"]
for t in [
    "Step 1. Expected one-year rates from the market: Kim-Wright (Federal Reserve Board) splits each zero-coupon "
    "Treasury yield, 1-10 years, into the expected average short rate and a term premium (FRED THREEFYn minus "
    "THREEFYTPn). The expected one-year rate in year n = n x a(n) - (n-1) x a(n-1); after year 10 it is held at the "
    f"year-10 rate. Today: y1 {y30[0]:.2f}, y2 {y30[1]:.2f}, y3 {y30[2]:.2f}, y4 {y30[3]:.2f}, y5 {y30[4]:.2f}, "
    f"y10 onward {y30[9]:.2f}.",
    "Step 2. n-year Treasury each month = [(1+y1)(1+y2)...(1+yn)]^(1/n) - 1 + maturity premium, n = 30 and 10. "
    f"Each premium is today's yield minus today's compounded expectation, held constant: 30y "
    f"{mpi['treasury_30y_now']:.2f} - {mpi['expected_30y_now']:.2f} = {mpi['maturity_premium_30y']:.2f}pp; 10y "
    f"{mpi['treasury_10y_now']:.2f} - {mpi['expected_10y_now']:.2f} = {mpi['maturity_premium_10y']:.2f}pp "
    f"(Kim-Wright's own 10y term premium: {mpi['kim_wright_10y_term_premium']:.2f}pp). Each path starts at today's "
    "yield and rolls as near-term years leave the window.",
    f"Long run: 30y = {y30[-1]:.2f} + {mpi['maturity_premium_30y']:.2f} = {mpi['treasury_30y_long_run']:.2f}%; "
    f"10y = {y30[-1]:.2f} + {mpi['maturity_premium_10y']:.2f} = {mpi['treasury_10y_long_run']:.2f}%.",
    f"Step 3. Mortgage = 30y Treasury + {mpi['spread_median_1977_2026']:.2f}pp, the median weekly PMMS minus 30y Treasury "
    f"spread since {mpi['spread_first']} ({mpi['spread_n_weeks']:,} weeks; mean {mpi['spread_mean']:.2f}, latest "
    f"{mpi['spread_latest']:.2f}).",
    "2026 = actual averages to date (Treasuries, PMMS, fed funds) blended with the path for the remaining weeks.",
]:
    r = note(ws, r, t)
r += 1
rows = [[n, k.fitted_yield, k.term_premium, k.expected_avg_short, k.expected_1y_rate,
         mpdi["y1_to_y10_now"][n - 1]] for n, k in kw.iterrows()]
r = table(ws, r, "Market vs FOMC dots: expected one-year rates from today",
          ["Year from now", "Kim-Wright fitted yield %", "Kim-Wright term premium (pp)",
           "Expected average short rate %", "Market: expected 1y rate in that year %", "FOMC dots: expected 1y rate %"],
          rows, [None, "0.00", "0.00", "0.00", "0.00", "0.00"])
rows = [[yr, r.treasury_30y, mpd.at[yr, "treasury_30y"], r.treasury_10y, mpd.at[yr, "treasury_10y"],
         r.mortgage_30y, mpd.at[yr, "mortgage_30y"]] for yr, r in mpa.iterrows()]
r = table(ws, r, "Market (BASE) vs FOMC dots alternative",
          ["Year", "30y: market %", "30y: dots %", "10y: market %", "10y: dots %", "Mortgage: market %",
           "Mortgage: dots %"], rows, [None, "0.00", "0.00", "0.00", "0.00", "0.00", "0.00"])
for t in [
    "FOMC dots alternative (16-Sep-2026 median: end-2026 4.1, 2027 4.1, 2028 3.9, 2029 3.6, longer run 3.2 by end-2030): "
    f"premiums 30y {mpdi['maturity_premium_30y']:.2f}pp, 10y {mpdi['maturity_premium_10y']:.2f}pp. The dots sit about "
    "0.8pp below market expectations from year 5 on, mostly the 3.2% longer-run rate against about 4.0% priced; "
    "backed out of today's yields, that gap would sit inside the premium, which is why the market path is BASE.",
    "Sources: FRED THREEFY1-10, THREEFYTP1-10 (Kim-Wright), DGS30, DGS10, DFF, MORTGAGE30US; "
    "federalreserve.gov/monetarypolicy/fomcprojtabl20260916.htm; FRED graph MORTGAGE30US minus DGS30 "
    "(data/raw/fred_MORTGAGE30US_minus_DGS30.csv).",
]:
    r = note(ws, r, t)
widths(ws, 12, 18, 10)

# 7. notes
ws = sheet(wb, "7 Notes", "What is behind these numbers", "Full method and diagnostics: docs/ in the repository.")
lines = [
    "Model: yearly flow into ownership within each 5-year age cohort = age effect + last year's ln(price/income), "
    "ln(mortgage payment factor) and delinquency, OLS 1992-2025 (n=352), year-clustered errors.",
    "Price and rate act on ages 25-54 only; 55+ respond to delinquency only. Adjusted R2 0.793.",
    "Forecast: fitted flows rolled forward through the cohorts; under-25 ownership held at 2025.",
    "Households: Census 2023 national population projections by age (low-immigration series for BASE), 2025 headship "
    "by age held, calibrated to JCHS 2025-35 change by ten-year age group (690k a year in total).",
    f"Price/income BASE: Fannie Mae HPI forecast (11-Sep-2026: +2.3% 2026, +1.0% 2027) against NAR family income at "
    f"its 2015-25 trend (+4.5% a year); flat after 2027. 2025 value {assum['price_to_income_2025']:.3f}.",
    "BASE mortgage rate: market expected short rates (Kim-Wright) compounded into a 30y Treasury path plus the "
    "median mortgage spread (sheet 6); the FOMC dots path is a scenario; "
    "flat 6-8% rows are sensitivities.",
    f"Delinquency held at {assum['delinq_now']:.2f}% (latest four quarters). NAR rate = PMMS + "
    f"{assum['nar_minus_pmms_spread']:.2f}pp (2021-25 average).",
    "Known bias: over 2022-25 the model (estimated to 2021) was 1.5pp too low on the national rate because it has no "
    "mortgage lock-in term; the near-term path for ages 35-54 is more likely too low than too high.",
    "2025 levels are Census CPS/HVS annual households (86.66m owners); the quarterly owner-occupied series averages "
    "86.71m for 2025. Use growth rates when splicing onto that series.",
    "Sources: Census HVS Tables 12/12a; NAR affordability data via HUD USHMC HousingDemand.xlsx; Freddie Mac PMMS; "
    "Federal Reserve delinquency (FRED DRSFRMACBS); Census 2023 population projections; JCHS; Fannie Mae ESR.",
]
for i, t in enumerate(lines, start=4):
    c = ws.cell(i, 1, t)
    c.font, c.alignment = f(), Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.freeze_panes = None
wb.save(OUT)
print("saved", OUT)
